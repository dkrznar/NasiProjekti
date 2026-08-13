from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from datetime import datetime
from models import db, Projekt, Biljeska, STATUSI, Aneks, ULOGE
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pomocno import parsiraj_datum, admin_required
from flask_login import LoginManager, login_user, logout_user, login_required, current_user


app = Flask(__name__)
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///projekti.db"
app.config["SECRET_KEY"] = "zabok123"

db.init_app(app)
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "prijava" 

from models import Korisnik

@login_manager.user_loader
def ucitaj_korisnika(korisnik_id):
    return Korisnik.query.get(int(korisnik_id))

from socijala import socijala
app.register_blueprint(socijala)

with app.app_context():
    db.create_all()

@app.template_filter("eur")
def eur(vrijednost):
    if vrijednost is None:
        vrijednost = 0
    s = f"{vrijednost:,.2f}"
    s = s.replace(",", "X").replace(".",",").replace("X", ".")
    return s + "\u00a0€"

@app.route("/")
def pocetna():
    return render_template("pocetna.html")

@app.route("/projekti")
def index():
    upit = Projekt.query

    # Filter: status
    status = request.args.get("status", "")
    if status:
        upit = upit.filter(Projekt.status == status)

    # Filter: datum početka od-do
    od = parsiraj_datum(request.args.get("od"))
    do_ = parsiraj_datum(request.args.get("do"))
    if od:
        upit = upit.filter(Projekt.datum_pocetka >= od)
    if do_:
        upit = upit.filter(Projekt.datum_pocetka <= do_)

    # Sortiranje
    sort = request.args.get("sort", "pocetak_novi")
    if sort == "abeceda":
        upit = upit.order_by(Projekt.naziv.asc())
    elif sort == "pocetak_stari":
        upit = upit.order_by(Projekt.datum_pocetka.asc())
    elif sort == "zavrsetak":
        upit = upit.order_by(Projekt.datum_zavrsetka.asc())
    else:  # pocetak_novi (zadano)
        upit = upit.order_by(Projekt.datum_pocetka.desc())

    projekti = upit.all()
    return render_template("index.html", projekti=projekti, statusi=STATUSI)

@app.route("/projekt/novi", methods=["GET", "POST"])
def novi_projekt():
    if request.method == "POST":
        p = Projekt(
            naziv=request.form["naziv"].strip(),
            klasa=request.form.get("klasa", "").strip(),
            opis=request.form.get("opis", "").strip(),
            datum_pocetka=parsiraj_datum(request.form.get("datum_pocetka")),
            datum_zavrsetka=parsiraj_datum(request.form.get("datum_zavrsetka")),
            status=request.form.get("status", "U pripremi"),
            naziv_poziva=request.form.get("naziv_poziva", "").strip(),
            nadlezno_tijelo=request.form.get("nadlezno_tijelo", "").strip(),
            dobiveni_iznos=float(request.form.get("dobiveni_iznos") or 0),
            nas_iznos=float(request.form.get("nas_iznos") or 0),
            datum_ugovora_financiranje=parsiraj_datum(request.form.get("datum_ugovora_financiranje")),
            datum_ugovora_nabava=parsiraj_datum(request.form.get("datum_ugovora_nabava")),
            izvodac=request.form.get("izvodac", "").strip(),
            rok_izvrsenja=parsiraj_datum(request.form.get("rok_izvrsenja")),
            unio=current_user.ime_prezime,
        )
        db.session.add(p)
        db.session.commit()
        flash("Projekt je spremljen.", "success")
        return redirect(url_for("index"))
 
    return render_template("form.html", statusi=STATUSI, projekt = None)
 

 
@app.route("/projekt/<int:projekt_id>/uredi", methods=["GET", "POST"])
def uredi_projekt(projekt_id):
    projekt = Projekt.query.get_or_404(projekt_id)

    if request.method == "POST":
        projekt.naziv = request.form["naziv"].strip()
        projekt.klasa=request.form.get("klasa", "").strip()
        projekt.opis=request.form.get("opis", "").strip()
        projekt.datum_pocetka=parsiraj_datum(request.form.get("datum_pocetka"))
        projekt.datum_zavrsetka=parsiraj_datum(request.form.get("datum_zavrsetka"))
        projekt.status=request.form.get("status", "U pripremi")
        projekt.naziv_poziva=request.form.get("naziv_poziva", "").strip()
        projekt.nadlezno_tijelo=request.form.get("nadlezno_tijelo", "").strip()
        projekt.dobiveni_iznos=float(request.form.get("dobiveni_iznos") or 0)
        projekt.nas_iznos=float(request.form.get("nas_iznos") or 0)
        projekt.datum_ugovora_financiranje=parsiraj_datum(request.form.get("datum_ugovora_financiranje"))
        projekt.datum_ugovora_nabava=parsiraj_datum(request.form.get("datum_ugovora_nabava"))
        projekt.izvodac=request.form.get("izvodac", "").strip()
        projekt.rok_izvrsenja=parsiraj_datum(request.form.get("rok_izvrsenja"))
        projekt.izmijenio=current_user.ime_prezime
        projekt.vrijeme_izmjene = datetime.now()

        db.session.commit()
        flash("Projekt je ažuriran.", "success")
        return redirect(url_for("projekt_detalji", projekt_id=projekt.id))
    
    return render_template("form.html", statusi= STATUSI, projekt=projekt)

@app.route("/projekt/<int:projekt_id>/obrisi", methods=["POST"])
def obrisi_projekt(projekt_id):
    projekt = Projekt.query.get_or_404(projekt_id)
    db.session.delete(projekt)
    db.session.commit()
    flash("Projekt je obrisan.", "warning")
    return redirect(url_for("index"))
 
@app.route("/projekt/<int:projekt_id>")
def projekt_detalji(projekt_id):
    projekt = Projekt.query.get_or_404(projekt_id)
    return render_template("projekt_detalji.html", projekt=projekt)

@app.route("/projekt/<int:projekt_id>/aneks", methods=["POST"])
def novi_aneks(projekt_id):
    projekt = Projekt.query.get_or_404(projekt_id)
    a = Aneks(
        projekt_id=projekt.id,
        datum=parsiraj_datum(request.form.get("datum")),
        napomena=request.form.get("napomena", "").strip(),
        dodao=current_user.ime_prezime,
    )
    db.session.add(a)
    db.session.commit()
    flash("Aneks je dodan.", "success")
    return redirect(url_for("projekt_detalji", projekt_id=projekt.id))

@app.route("/aneks/<int:aneks_id>/obrisi", methods=["POST"])
def obrisi_aneks(aneks_id):
    a = Aneks.query.get_or_404(aneks_id)
    projekt_id = a.projekt_id
    db.session.delete(a)
    db.session.commit()
    flash("Aneks je obrisan.", "warning")
    return redirect(url_for("projekt_detalji", projekt_id=projekt_id))

@app.route("/projekt/<int:projekt_id>/biljeska", methods=["POST"])
def nova_biljeska(projekt_id):
    projekt = Projekt.query.get_or_404(projekt_id)
    b = Biljeska(
        projekt_id=projekt.id,
        tekst=request.form["tekst"].strip(),
        ime=current_user.ime_prezime,
    )
    db.session.add(b)
    db.session.commit()
    flash("Bilješka je dodana.", "success")
    return redirect(url_for("projekt_detalji", projekt_id=projekt.id))

@app.route("/biljeska/<int:biljeska_id>/uredi", methods=["GET", "POST"])
def uredi_biljesku(biljeska_id):
    b = Biljeska.query.get_or_404(biljeska_id)
    if request.method == "POST":
        b.tekst = request.form["tekst"].strip()
        db.session.commit()
        flash("Bilješka je ažurirana.", "success")
        return redirect(url_for("projekt_detalji", projekt_id=b.projekt_id))
    return render_template("uredi_biljeska.html", biljeska=b)

@app.route("/biljeska/<int:biljeska_id>/obrisi", methods=["POST"])
def obrisi_biljesku(biljeska_id):
    b = Biljeska.query.get_or_404(biljeska_id)
    projekt_id = b.projekt_id
    db.session.delete(b)
    db.session.commit()
    flash("Bilješka je obrisana.", "warning")
    return redirect(url_for("projekt_detalji", projekt_id=projekt_id))

@app.route("/export")
def export_excel():
    wb = Workbook()

    # --- Sheet 1: Projekti ---
    ws = wb.active
    ws.title = "Projekti"

    zaglavlje = ["Naziv projekta", "Klasa", "Naziv poziva", "Nadležno tijelo",
                 "Datum početka", "Datum završetka", "Status",
                 "Dobiveni iznos (€)", "Vlastiti iznos (€)", "Ukupan iznos (€)",
                 "Ugovor o financiranju", "Ugovor o nabavi", "Izvođač",
                 "Aneksi", "Rok izvršenja", "Opis", "Unio/la"]
    ws.append(zaglavlje)
    for celija in ws[1]:
        celija.font = Font(bold=True)

    def dat(d):
        return d.strftime("%d.%m.%Y.") if d else ""

    projekti = Projekt.query.order_by(Projekt.datum_pocetka.desc()).all()
    for p in projekti:
        ws.append([
            p.naziv, p.klasa, p.naziv_poziva, p.nadlezno_tijelo,
            dat(p.datum_pocetka), dat(p.datum_zavrsetka), p.status,
            p.dobiveni_iznos or 0, p.nas_iznos or 0, p.ukupan_iznos,
            dat(p.datum_ugovora_financiranje), dat(p.datum_ugovora_nabava),
            p.izvodac, ", ".join(a.datum.strftime("%d.%m.%Y.") for a in p.aneksi), dat(p.rok_izvrsenja),
            p.opis, p.unio,
        ])

    sirine = [30, 18, 25, 25, 13, 13, 12, 15, 15, 15, 15, 15, 25, 13, 13, 40, 15]
    for i, s in enumerate(sirine, start=1):
        ws.column_dimensions[get_column_letter(i)].width = s

    # --- Sheet 2: Bilješke ---
    ws2 = wb.create_sheet("Bilješke")
    ws2.append(["Projekt", "Bilješka", "Napisao/la", "Vrijeme"])
    for celija in ws2[1]:
        celija.font = Font(bold=True)

    for p in projekti:
        for b in p.biljeske:
            ws2.append([
                p.naziv, b.tekst, b.ime,
                b.vrijeme.strftime("%d.%m.%Y. %H:%M") if b.vrijeme else "",
            ])

    for i, s in enumerate([30, 60, 15, 18], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = s

    # --- Pošalji datoteku ---
    datoteka = BytesIO()
    wb.save(datoteka)
    datoteka.seek(0)

    naziv = f"projekti_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    return send_file(
        datoteka,
        as_attachment=True,
        download_name=naziv,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

@app.route("/ime", methods=["GET", "POST"])
def postavi_ime():
    if request.method == "POST":
        ime = request.form["ime"].strip()
        odgovor = redirect(request.form.get("next") or url_for("index"))
        odgovor.set_cookie("korisnik", ime, max_age=60*60*24*365)
        return odgovor
    return render_template("ime.html", next=request.args.get("next", ""))

@app.before_request
def provjeri_prijavu():
    if request.endpoint in ("prijava", "static"):
        return
    if not current_user.is_authenticated:
        return redirect(url_for("prijava", next=request.url))
    if request.endpoint not in ("promijeni_svoju_lozinku", "odjava") and current_user.mora_promijeniti_lozinku:
        return redirect(url_for("promijeni_svoju_lozinku"))

@app.route("/prijava", methods=["GET", "POST"])
def prijava():
    if request.method == "POST":
        korisnicko_ime = request.form["korisnicko_ime"].strip()
        lozinka = request.form["lozinka"]

        korisnik = Korisnik.query.filter_by(korisnicko_ime=korisnicko_ime).first()

        if korisnik and korisnik.aktivan and korisnik.provjeri_lozinku(lozinka):
            login_user(korisnik)
            if korisnik.mora_promijeniti_lozinku:
                flash("Molimo postavite novu lozinku.", "warning")
                return redirect(url_for("promijeni_svoju_lozinku"))
            flash(f"Dobrodošla/o, {korisnik.ime_prezime}!", "success")
            return redirect(request.args.get("next") or url_for("pocetna"))
        else:
            flash("Pogrešno korisničko ime ili lozinka.", "danger")

    return render_template("prijava.html")


@app.route("/odjava")
@login_required
def odjava():
    logout_user()
    flash("Uspješno ste se odjavili.", "success")
    return redirect(url_for("prijava"))

@app.route("/korisnici")
@admin_required
def korisnici():
    popis = Korisnik.query.order_by(Korisnik.korisnicko_ime.asc()).all()
    return render_template("korisnici.html", korisnici=popis, uloge=ULOGE)


@app.route("/korisnici/novi", methods=["POST"])
@admin_required
def novi_korisnik():
    if Korisnik.query.filter_by(korisnicko_ime=request.form["korisnicko_ime"].strip()).first():
        flash("Korisničko ime već postoji.", "danger")
        return redirect(url_for("korisnici"))

    k = Korisnik(
        korisnicko_ime=request.form["korisnicko_ime"].strip(),
        ime_prezime=request.form["ime_prezime"].strip(),
        uloga=request.form.get("uloga", "urednik"),
    )
    k.postavi_lozinku(request.form["lozinka"])
    db.session.add(k)
    db.session.commit()
    flash(f"Korisnik {k.ime_prezime} je dodan.", "success")
    return redirect(url_for("korisnici"))


@app.route("/korisnici/<int:korisnik_id>/deaktiviraj", methods=["POST"])
@admin_required
def deaktiviraj_korisnika(korisnik_id):
    k = Korisnik.query.get_or_404(korisnik_id)
    k.aktivan = not k.aktivan   # preklopnik: uključi/isključi
    db.session.commit()
    flash(f"Korisnik {k.ime_prezime} je {'aktiviran' if k.aktivan else 'deaktiviran'}.", "success")
    return redirect(url_for("korisnici"))


@app.route("/korisnici/<int:korisnik_id>/reset-lozinke", methods=["POST"])
@admin_required
def reset_lozinke(korisnik_id):
    k = Korisnik.query.get_or_404(korisnik_id)
    k.postavi_lozinku(request.form["nova_lozinka"])
    k.mora_promijeniti_lozinku = True
    db.session.commit()
    flash(f"Lozinka za {k.ime_prezime} je promijenjena.", "success")
    return redirect(url_for("korisnici"))

@app.route("/korisnici/<int:korisnik_id>/promijeni-ulogu", methods=["POST"])
@admin_required
def promijeni_ulogu(korisnik_id):
    k = Korisnik.query.get_or_404(korisnik_id)
    k.uloga = request.form.get("uloga", k.uloga)
    db.session.commit()
    flash(f"Uloga za {k.ime_prezime} je promijenjena u '{k.uloga}'.", "success")
    return redirect(url_for("korisnici"))


@app.route("/moja-lozinka", methods=["GET", "POST"])
@login_required
def promijeni_svoju_lozinku():
    if request.method == "POST":
        nova = request.form["nova_lozinka"]
        ponovljena = request.form["ponovljena_lozinka"]

        if nova != ponovljena:
            flash("Lozinke se ne podudaraju.", "danger")
            return redirect(url_for("promijeni_svoju_lozinku"))

        if len(nova) < 6:
            flash("Lozinka mora imati barem 6 znakova.", "danger")
            return redirect(url_for("promijeni_svoju_lozinku"))

        current_user.postavi_lozinku(nova)
        current_user.mora_promijeniti_lozinku = False
        db.session.commit()
        flash("Lozinka je uspješno promijenjena.", "success")
        return redirect(url_for("pocetna"))

    return render_template("promijeni_lozinku.html")

if __name__ == "__main__":
    from waitress import serve
    print("Aplikacija radi na http://192.168.0.13:5000")
    serve(app, host="0.0.0.0", port=5000)